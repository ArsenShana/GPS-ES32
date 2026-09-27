from pathlib import Path
import textwrap
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).parent
W, H = 960, 540
INK, MUTED, BLUE, LIGHT, LINE = '172333', '607083', '2563EB', 'F2F6FC', 'E2E8F0'
FONT = '/System/Library/Fonts/Supplemental/Arial.ttf'
BOLD = '/System/Library/Fonts/Supplemental/Arial Bold.ttf'
pdfmetrics.registerFont(TTFont('Arial', FONT))
pdfmetrics.registerFont(TTFont('Arial-Bold', BOLD))
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
pdf = canvas.Canvas(str(ROOT/'GPS_ideas.pdf'), pagesize=(W,H))
pdf.setTitle('GPS: варианты интеграции — AlanaTech')
previews=[]

def color(c): return RGBColor.from_string(c)
def rect(x,y,w,h,fill=LIGHT):
    s=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Pt(x),Pt(y),Pt(w),Pt(h))
    s.fill.solid(); s.fill.fore_color.rgb=color(fill); s.line.fill.background()
    pdf.setFillColor('#'+fill); pdf.rect(x,H-y-h,w,h,fill=1,stroke=0)
    draw.rectangle((x*2,y*2,(x+w)*2,(y+h)*2),fill='#'+fill)

def text(x,y,w,content,size=18,bold=False,c=INK,leading=1.24):
    font='Arial-Bold' if bold else 'Arial'
    lines=[]
    for para in content.split('\n'):
        line=''
        for word in para.split():
            trial=(line+' '+word).strip()
            if pdfmetrics.stringWidth(trial,font,size)>w and line:
                lines.append(line); line=word
            else: line=trial
        lines.append(line)
    box=slide.shapes.add_textbox(Pt(x),Pt(y),Pt(w+4),Pt(len(lines)*size*leading+5))
    tf=box.text_frame; tf.clear(); tf.word_wrap=False
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name='Arial'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=color(c)
        p.space_before=Pt(0); p.space_after=Pt(0); p.line_spacing=Pt(size*leading)
        yy=y+i*size*leading
        pdf.setFont(font,size); pdf.setFillColor('#'+c); pdf.drawString(x,H-yy-size*.94,line)
        draw.text((x*2,yy*2),line,font=ImageFont.truetype(BOLD if bold else FONT,round(size*2)),fill='#'+c)
    assert y+len(lines)*size*leading < H-7, (content,y)

def picture(name,x,y,w,h):
    path=ROOT/'assets'/name
    im=Image.open(path).convert('RGB'); iw,ih=im.size
    scale=min(w/iw,h/ih); pw,ph=iw*scale,ih*scale
    xx,yy=x+(w-pw)/2,y+(h-ph)/2
    slide.shapes.add_picture(str(path),Pt(xx),Pt(yy),width=Pt(pw),height=Pt(ph))
    pdf.drawImage(str(path),xx,H-yy-ph,pw,ph)
    preview.paste(im.resize((round(pw*2),round(ph*2))), (round(xx*2),round(yy*2)))

def new(n,kicker,title=None,subtitle=None):
    global slide, preview, draw
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb=color('FFFFFF')
    preview=Image.new('RGB',(W*2,H*2),'white'); draw=ImageDraw.Draw(preview)
    text(42,24,800,kicker.upper(),10,True,BLUE)
    if title: text(42,59,890,title,30,True)
    if subtitle: text(42,107,880,subtitle,14,False,MUTED)
    rect(42,505,876,1,LINE)
    text(42,516,800,'ALANATECH  /  GPS • концепции для обсуждения',9,False,MUTED)
    text(890,515,30,f'{n:02}',10,True,MUTED)

def done():
    previews.append(preview.copy()); pdf.showPage()

new(1,'Продуктовые идеи / для тимлида')
text(42,74,880,'Куда встроить GPS',44,True)
text(44,136,850,'Три основных варианта и ещё три направления для развития',20,False,MUTED)
for x,name,label,sub in [(42,'powerbank.png','Пауэрбанк','Питание + поиск устройства'),(340,'bicycle.png','Велосипед','Местоположение + движение'),(638,'dashcam.png','Камера в машине','Видео + маршрут поездки')]:
    picture(name,x,208,280,192)
    text(x+7,418,275,label,21,True)
    text(x+7,451,275,sub,13,False,MUTED)
text(44,183,850,'Иллюстрации — концепты возможной интеграции, не готовые изделия.',11,False,MUTED)
done()

def option(n,idx,title,sub,img,blocks,bottom):
    new(n,f'Основной вариант {idx}',title,sub)
    picture(img,22,161,518,316)
    text(46,478,465,'Концепт • размещение модуля показано условно',10,False,MUTED)
    y=162
    for label,body in blocks:
        text(567,y,344,label,11,True,BLUE)
        text(567,y+23,344,body,17)
        y+=98
    rect(556,458,361,33)
    text(567,464,339,bottom,11,True,BLUE)
    done()

option(2,'01','GPS в пауэрбанке','Зарядка телефона и поиск самого устройства в одном корпусе.','powerbank.png',[
    ('ГДЕ РАЗМЕСТИТЬ','В пластиковом корпусе: модуль и антенна рядом с платой питания.'),
    ('ЗАЧЕМ','Показать на карте пауэрбанк или сумку, в которой он находится.'),
    ('ЧТО ПРОВЕРИТЬ','Расход заряда, нагрев и качество приёма рядом с аккумулятором.')
],'Кандидат для первого прототипа')
option(3,'02','GPS в велосипеде','Поиск велосипеда и уведомление о начале движения.','bicycle.png',[
    ('ГДЕ РАЗМЕСТИТЬ','В заднем фонаре или отдельном пластиковом блоке под седлом.'),
    ('ЗАЧЕМ','Видеть последнюю позицию; получать сигнал о движении при наличии связи.'),
    ('ЧТО ПРОВЕРИТЬ','Защиту от воды, крепление, зарядку и приём сигнала рядом с рамой.')
],'Для уведомлений нужен датчик движения')
option(4,'03','GPS в автомобильной камере','Видеорегистратор, который связывает видео с маршрутом поездки.','dashcam.png',[
    ('ГДЕ РАЗМЕСТИТЬ','В корпусе регистратора или его креплении у лобового стекла.'),
    ('ЗАЧЕМ','Записывать координаты, время и скорость вместе с видео.'),
    ('ЧТО ПРОВЕРИТЬ','Совместимость с камерой, нагрев на солнце и питание на парковке.')
],'Удалённый просмотр требует канала связи')

new(5,'Дополнительные направления','Ещё три варианта','Идеи для следующего этапа — после проверки основного прототипа.')
cards=[('04','Рюкзак / сумка','Куда: съёмный модуль во внутреннем кармане.','Польза: поиск своей вещи.','Проверить: зарядку и приём внутри здания.'),('05','Кейс инструмента','Куда: отдельный блок под пластиковой крышкой.','Польза: местоположение рабочего комплекта.','Проверить: экранирование металлическими предметами.'),('06','Электросамокат','Куда: корпус панели или пластиковый отсек.','Польза: положение и история поездок.','Проверить: питание, вибрации и влагозащиту.')]
for i,(num,title,a,b,c) in enumerate(cards):
    x=42+i*299
    rect(x,163,278,314,LIGHT)
    text(x+20,181,235,num,34,True,BLUE)
    text(x+20,240,242,title,20,True)
    text(x+20,284,238,a,16)
    text(x+20,352,238,b,16)
    text(x+20,413,238,c,14,False,MUTED)
done()

new(6,'Принцип работы','Как координаты попадут на телефон','GPS определяет положение. Передача координат — отдельная часть системы.')
for x,num,title,body in [(42,'1','GPS / GNSS','Получает сигналы спутников и определяет координаты.'),(341,'2','Канал связи','Сотовый модем отправляет данные при наличии покрытия.'),(640,'3','Приложение','Показывает позицию, время обновления и историю.')]:
    rect(x,176,278,198)
    text(x+20,192,235,num,28,True,BLUE)
    text(x+20,240,239,title,22,True)
    text(x+20,282,237,body,17)
text(321,256,24,'›',25,True,BLUE); text(620,256,24,'›',25,True,BLUE)
text(45,410,864,'Без модема: можно хранить маршрут в устройстве и выгружать его позже.',17,True)
text(45,452,864,'Bluetooth подходит для связи рядом с телефоном; для удалённого поиска нужен другой канал.',14,False,MUTED)
slide.notes_slide.notes_text_frame.text='Схема упрощена. Для удалённого доступа также нужна серверная часть. Доступ к координатам предоставляется владельцу устройства. При потере связи следует отображать время последнего обновления.'
done()

new(7,'Сравнение','Что выбрать для первого прототипа','Предварительная оценка концепций; стоимость и автономность ещё предстоит измерить.')
cols=[42,245,474,704]; widths=[203,229,230,214]
headers=['Критерий','Пауэрбанк','Велосипед','Автокамера']
for x,w,t in zip(cols,widths,headers):
    rect(x,162,w,44,BLUE); text(x+12,173,w-23,t,16,True,'FFFFFF')
rows=[['Главная польза','Поиск устройства','Поиск + движение','Видео + маршрут'],['Питание','Своя батарея','Отдельная батарея','Бортовая сеть'],['Главная задача','Расход заряда','Вода и крепление','Связь с камерой'],['Приоритет идеи','Первый прототип','Следующий вариант','Отдельное направление']]
for j,row in enumerate(rows):
    for x,w,t in zip(cols,widths,row):
        rect(x,206+j*50,w,49,LIGHT if j%2==0 else 'FFFFFF')
        text(x+12,219+j*50,w-22,t,14,j==0 and x==42)
rect(42,432,876,55,LIGHT)
text(58,442,845,'Предлагаемый старт: пауэрбанк — готовый запас энергии и место для электроники.',16,True,BLUE)
done()

new(8,'Решение для обсуждения','Начать с GPS-пауэрбанка','Рабочая гипотеза: проверить базовый трекер, затем адаптировать его к другим корпусам.')
for y,num,title,body in [(163,'01','Собрать макет','GPS-модуль, контроллер, питание и выбранный канал связи.'),(256,'02','Проверить в реальных условиях','Приём на улице и в сумке, расход энергии, передачу координат.'),(349,'03','Показать результат','Позиция на карте, история маршрута и время последнего обновления.')]:
    rect(42,y,54,54,LIGHT); text(53,y+10,43,num,23,True,BLUE)
    text(116,y,790,title,22,True); text(116,y+36,790,body,17,False,MUTED)
rect(42,451,876,38,BLUE)
text(57,459,845,'Обсудить с тимлидом: сценарий поиска, размеры корпуса и желаемую автономность.',15,True,'FFFFFF')
done()

prs.core_properties.title='Куда встроить GPS'
prs.core_properties.subject='Варианты интеграции для обсуждения с тимлидом'
prs.core_properties.author='AlanaTech'
prs.save(ROOT/'GPS_ideas.pptx'); pdf.save()
contact=Image.new('RGB',(960,1080),'#E2E8F0')
for i,im in enumerate(previews):
    im.save(ROOT/f'preview-{i+1:02}.png')
    contact.paste(im.resize((480,270)),((i%2)*480,(i//2)*270))
contact.save(ROOT/'overview.jpg')
print('Created PPTX, PDF and previews:',len(prs.slides),'slides')
